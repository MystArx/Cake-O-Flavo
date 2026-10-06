import io
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from core import load_all

f = lambda v: float(v or 0)
HEAD = PatternFill("solid", fgColor="3B0F26")

def totals(i):
    sub = sum(f(t.get("qty")) * f(t.get("rate")) for t in i.get("items", []))
    x = max(0, sub - f(i.get("disc"))); tax = x * f(i.get("tax")) / 100
    return sub, f(i.get("disc")), tax, x + tax

def day(s):
    try: return datetime.strptime(s, "%Y-%m-%d").date()
    except Exception: return None

def sheet(wb, title, headers, rows, widths, fmts=None):
    ws = wb.create_sheet(title); ws.append(headers)
    for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = HEAD; c.alignment = Alignment(vertical="center")
    for r in rows: ws.append(r)
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    for col, fmt in (fmts or {}).items():
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            for c in row: c.number_format = fmt
    ws.freeze_panes = "A2"
    if rows: ws.auto_filter.ref = ws.dimensions
    return ws

def make_xlsx(frm="", to=""):
    d = load_all(); st = d["set"] or {}; cur = st.get("cur", "₹")
    inr = lambda s: (not frm or s >= frm) and (not to or s <= to)
    I = sorted([i for i in d["inv"] if inr(i.get("date", ""))], key=lambda i: (i.get("date", ""), i.get("no", 0)))
    X = sorted([e for e in d["exp"] if inr(e.get("date", ""))], key=lambda e: e.get("date", ""))
    money, DATE = "#,##0.00", "DD MMM YYYY"
    paid = lambda i: i.get("status") == "paid"
    sales = sum(totals(i)[3] for i in I); coll = sum(totals(i)[3] for i in I if paid(i)); spent = sum(f(e.get("amt")) for e in X)
    wb = Workbook(); ws = wb.active; ws.title = "Summary"
    rows = [("Business", st.get("name", "")), ("Period", f"{frm} to {to}" if frm else "All time"), ("Invoices", len(I)),
            ("Sales (billed)", sales), ("Collected (paid)", coll), ("Still to collect", sales - coll), ("Expenses", spent),
            ("Profit (collected - expenses)", coll - spent), ("Average invoice", sales / len(I) if I else 0),
            ("To collect, all time", sum(totals(i)[3] for i in d["inv"] if not paid(i)))]
    for r in rows: ws.append(r)
    ws.column_dimensions["A"].width = 32; ws.column_dimensions["B"].width = 26
    for r in ws.iter_rows(): r[0].font = Font(bold=True)
    for r in ws.iter_rows(min_row=4): r[1].number_format = money
    ws["B9"].number_format = money
    ws["B3"].number_format = "0"
    sheet(wb, "Invoices", ["Invoice", "Date", "Customer", "Subject", "Status", "Sub total", "Discount", "Tax", "Total"],
          [["INV-%06d" % int(i.get("no", 0)), day(i.get("date", "")), i.get("cname", ""), i.get("subject", ""), "Paid" if paid(i) else "Unpaid", *totals(i)] for i in I],
          [14, 14, 28, 34, 10, 14, 12, 12, 14], {2: DATE, 6: money, 7: money, 8: money, 9: money})
    items = [["INV-%06d" % int(i.get("no", 0)), day(i.get("date", "")), i.get("cname", ""), t.get("name", ""), f(t.get("qty")), f(t.get("rate")), f(t.get("qty")) * f(t.get("rate"))] for i in I for t in i.get("items", [])]
    sheet(wb, "Items", ["Invoice", "Date", "Customer", "Item", "Qty", "Rate", "Amount"], items, [14, 14, 28, 34, 10, 14, 14], {2: DATE, 6: money, 7: money})
    sheet(wb, "Expenses", ["Date", "Category", "Amount", "Note"], [[day(e.get("date", "")), e.get("cat", ""), f(e.get("amt")), e.get("note", "")] for e in X], [14, 22, 14, 40], {1: DATE, 3: money})
    cr = []
    for c in d["cust"]:
        mine = [i for i in I if i.get("cname") == c["name"]]
        cr.append([c["name"], c.get("phone", ""), c.get("addr", ""), len(mine), sum(totals(i)[3] for i in mine), sum(totals(i)[3] for i in mine if paid(i)), sum(totals(i)[3] for i in mine if not paid(i))])
    sheet(wb, "Customers", ["Name", "Phone", "Address", "Invoices", "Billed", "Paid", "Unpaid"], cr, [28, 18, 34, 10, 14, 14, 14], {5: money, 6: money, 7: money})
    pr = {}
    for i in I:
        for t in i.get("items", []):
            v = pr.setdefault(t.get("name", ""), [0, 0]); v[0] += f(t.get("qty")); v[1] += f(t.get("qty")) * f(t.get("rate"))
    sheet(wb, "Products", ["Product", "Default rate", "Units sold", "Revenue"],
          [[p["name"], f(p.get("rate")), pr.get(p["name"], [0, 0])[0], pr.get(p["name"], [0, 0])[1]] for p in d["prod"]], [32, 14, 12, 14], {2: money, 4: money})
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()
